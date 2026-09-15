# -*- coding: utf-8 -*-
"""첫구매 실적 대시보드 데이터 빌드

사용법
  python build_data.py               # 이 폴더 아래 CSV가 든 모든 하위 폴더를 합쳐 data/data.js 생성
  python build_data.py 9월2주차       # 특정 폴더만

raw 규칙 (태블로 crosstab 내보내기 · UTF-16 탭 구분 CSV)
  파일명 키워드로 자동 인식한다.
    전체관점 … _첫구매 (일자별 / 주별) → 첫구매 일평균거래액·일평균고객수 (회원구분 × 채널)
    전체관점 파일명에 '첫구매'가 없으면 MALL 전체 거래액 → 현재 대시보드에서는 읽지 않음
      단, 파일명에 '첫구매'가 없어도 내용이 첫구매(기존회원 거래액 비중 70% 이하)면 첫구매로 읽고 경고를 찍는다
      (예: '전체관점 - 일자별 실적 (기본)_2025.csv' 는 주별 _첫구매 값과 정확히 일치하는 첫구매 파일이었다)
    비회원 트래픽 / 가입자수 / 당일가입 첫구매율 / 첫구매율   (채널)
    가입율 은 가입자수 ÷ 비회원트래픽으로 다시 계산하므로 읽지 않는다.
    조직 카테고리별 첫구매 실적 (xlsx) + 상품관점 … _첫구매 (CSV) → 상품 구성 드릴다운 data/products.js
    신규회원실적대시보드 … (CSV, 지표 × 연령대 × 채널 · 날짜 열 'YYYY. M. D.') → 총)/순) 첫구매 거래액·고객수 (당년신규) — KPI 실적
    회원현황 (CSV, 연도 없는 'M/D' + 요일 행) → 누적회원수 · 신규회원수 · 유효회원수 (앱 설치자 미포함 기준)
    PUSH (CSV, 연도 없는 'M/D') → 앱푸시 수신동의 회원 기존/신규/Total × 수신동의·증감·신규추가·이탈 (앱 설치자 기준)
    APP설치데이터 (xlsx, 2026-03-01~) → 전체·신규·재설치 · 스토어 방문 · 삭제 · Push 활성 기기 · 순증 설치
  폴더는 CSV 수정시각 순으로 읽고, 같은 날짜·주차 값은 나중 폴더가 덮어쓴다.
  하위 폴더까지 찾는다 — 예: 9월2주차(전체 이력) + 2026년/0914 · 2026년/0915 …(일자별 추가분)
  → 새 raw는 폴더째 넣고 이 스크립트만 다시 실행하면 된다.
  조직 카테고리별 첫구매 실적은 xlsx · CSV 둘 다 읽는다(일자별 추가분은 CSV).
  회원현황에서 세 값이 모두 0인 날은 아직 갱신 전으로 보고 건너뛴다.

백업
  빌드할 때마다 backup/첫구매대시보드_백업_데이터YYYYMMDD.json.gz 를 쓴다(실적 · KPI · 상품).
  대시보드 '백업 불러오기'로 열 수 있고, 거기에 '새 데이터 추가'로 일자별 raw 를 올려 합친 뒤 '백업 다운로드'하면
  원천 이력 폴더 없이도 다음 날 데이터를 이어서 갱신할 수 있다(합치는 규칙은 merge.js 가 이 파일과 같게 옮겨 둔 것).

집계 규칙
  첫구매       = 전체관점 _첫구매 파일의 *TOTAL (회원구분 1_당월신규 / 2_기가입신규 / 3_기존)
  주·월 실적   = 대시보드에서 일자 raw를 합산. 주별 raw는 일자 raw가 없는 기간(전년)의 보조값으로만 쓴다
  주차         = ISO 주(월~일). 라벨 'MM월 N주차' 는 그 주 목요일이 속한 달 기준
  주별 거래액·고객수 원천은 '일평균' 이므로 × 일수 로 주 합계로 바꿔 저장
  주별 비회원트래픽은 주간 중복제거 값(일자 합의 약 0.8배)이라 일자 합과 섞지 않고 tru / fpnu 로 따로 저장한다.
    태블로 주별 가입율·첫구매율의 분모가 이 주간 트래픽이다(가입율 88/88주 일치) → 대시보드 주별 보기는 이 값으로 계산
  비율(당일가입 CR·첫구매율)은 분모(가입자수·트래픽)를 곱해 분자로 저장 → 어떤 기간이든 재집계 가능
"""
import csv
import datetime as dt
import io
import json
import pickle
import re
import sys
from collections import Counter, defaultdict
from pathlib import Path

ROOT = Path(__file__).resolve().parent
OUT = ROOT / 'data' / 'data.js'
OUT_PROD = ROOT / 'data' / 'products.js'
CACHE_DIR = ROOT / 'data' / '_cache'
BACKUP_DIR = ROOT / 'backup'
KPI_JS = ROOT / 'kpi.js'
SKIP_DIRS = {'data', 'backup', '__pycache__', '.git', 'node_modules'}
BACKUP_FORMAT = 'fp-dashboard-backup'

CHANNEL_ORDER = ['*TOTAL', '직접', '광고', 'EP', 'PUSH', '제휴', '브랜드광고', '미디어커머스']
MEMBER = {'*TOTAL': 'T', '1_당월신규': '1', '2_기가입신규': '2', '3_기존': '3'}
OVERALL_METRIC = {'일평균거래액': 'amt', '일평균고객수': 'cust'}
# 앞에서부터 매칭 — '당일가입 첫구매율' 이 '첫구매율' 보다 먼저 와야 한다
SIMPLE_FILES = [('당일가입 첫구매율', 'cr'), ('첫구매율', 'fpr'),
                ('비회원 트래픽', 'tr'), ('비회원트래픽', 'tr'), ('가입자수', 'sg')]
RATE_TO_COUNT = [('cr', 'sg', 'crn'), ('fpr', 'tr', 'fpn')]

# 상품 구성 드릴다운 (채널 → BPU → 카테고리 → 브랜드 → 상품)
#   조직 카테고리별 첫구매 실적 xlsx : 상품 단위 행 (첫구매 거래액의 일부만 담김 → 화면에 커버리지 표기)
#   상품관점 일자별 실적 CSV        : 채널 × BPU 첫구매 거래액 실측(100%) → 커버리지 분모
PRODUCT_XLSX_KEY = '조직 카테고리별'
PRODUCT_CSV_KEY = '상품관점'
PRODUCT_TOP_N = 20  # 브랜드 경로(BPU·카테고리·브랜드)별 연간 상위 N 상품만 개별 보관, 나머지는 '기타 상품'
PRODUCT_COLS = {'date': '결제_일자', 'bpu': 'BPU', 'ch': 'AF대분류', 'cat': '대카테고리', 'brand': '브랜드',
                'code': '상품코드', 'name': '상품명', 'amt': '거래액', 'cust': '주문고객수'}

# 신규회원 실적 대시보드(첫구매) — '당년신규' 첫구매 실적 원천. 연령대 Total 행만 쓴다
NEWMEMBER_KEY = '신규회원실적대시보드'
NEWMEMBER_METRICS = {'총)첫구매 거래액 (당년신규)': 'nya', '총)첫구매 고객수 (당년신규)': 'nyc',
                     '순)첫구매 거래액 (당년신규)': 'nyna', '순)첫구매 고객수 (당년신규)': 'nync'}
# 앱 설치 · 앱푸시 수신동의 · 회원현황
APP_KEY = 'APP설치'
APP_COLS = {'전체 설치': 'all', '신규 설치': 'new', '재설치': 're', '스토어 방문': 'visit', '삭제': 'del',
            'Push 활성 기기': 'pushdev', '순증 설치': 'net'}
PUSH_KEY = 'PUSH'
PUSH_SEG = {'기존': 'ex', '신규': 'ny', 'Total': 'tot'}
PUSH_ITEM = {'수신동의': 'stock', '증감': 'chg', '신규추가(+)': 'add', '기존이탈(-)': 'out'}
MEMBER_STATUS_KEY = '회원현황'
MEMBER_STATUS = {'누적회원수(천)': 'cum', '신규회원수': 'new', '유효회원수': 'valid'}
CHANNEL_BASES = {'amt', 'cust', 'tr', 'sg', 'cr', 'fpr', 'nya', 'nyc', 'nyna', 'nync'}  # 키 뒤쪽이 채널명인 시리즈
WEEKDAY_KO = '월화수목금토일'

YEAR_RE = re.compile(r'^\d{4}$')
DATE_RE = re.compile(r'^(\d{1,2})/(\d{1,2})$')
WEEK_RE = re.compile(r'^(\d{1,2})월\s*(\d)주차$')


DRM_HEADER = b'SCDSA'  # 사내 문서보안(DRM)으로 암호화된 파일의 머리 바이트


def is_drm(path):
    with open(path, 'rb') as fh:
        return fh.read(len(DRM_HEADER)) == DRM_HEADER


def _raw_cache_file(path):
    return CACHE_DIR / 'raw' / f'{path.parent.name}__{path.name}.pkl'


def _save_raw(path, rows):
    """읽기에 성공한 raw 를 보관 — 나중에 사내 DRM 이 파일을 암호화해도 마지막 내용으로 빌드할 수 있게"""
    try:
        f = _raw_cache_file(path)
        f.parent.mkdir(parents=True, exist_ok=True)
        f.write_bytes(pickle.dumps((dt.datetime.now().strftime('%Y-%m-%d %H:%M'), rows)))
    except Exception:
        pass


def _drm_fallback(path):
    f = _raw_cache_file(path)
    if f.exists():
        saved, rows = pickle.loads(f.read_bytes())
        print(f'  [경고] {path.parent.name}/{path.name}: DRM 암호화 → {saved}에 읽어 둔 내용으로 대신 읽음')
        return rows
    raise ValueError('DRM 암호화 파일(SCDSA) — 태블로에서 다시 내보내거나 DRM 해제 후 넣어 주세요')


def read_rows(path):
    raw = path.read_bytes()
    if raw[:len(DRM_HEADER)] == DRM_HEADER:
        return _drm_fallback(path)
    if raw[:2] in (b'\xff\xfe', b'\xfe\xff'):
        text = raw.decode('utf-16')
    else:
        try:
            text = raw.decode('utf-8-sig')
        except UnicodeDecodeError:
            text = raw.decode('cp949')
    head = text.split('\n', 1)[0]
    delim = '\t' if head.count('\t') >= head.count(',') else ','
    rows = [[c.strip() for c in r] for r in csv.reader(io.StringIO(text), delimiter=delim)]
    _save_raw(path, rows)
    return rows


def to_num(s):
    s = s.replace(',', '')
    if s in ('', '-'):
        return None
    pct = s.endswith('%')
    try:
        v = float(s[:-1] if pct else s)
    except ValueError:
        return None
    return v / 100 if pct else v


_ISO_CACHE = {}


def iso_week_map(year):
    """(월, n주차) → (ISO 주번호, 월요일)"""
    if year not in _ISO_CACHE:
        out, seen = {}, Counter()
        for w in range(1, dt.date(year, 12, 28).isocalendar()[1] + 1):
            monday = dt.date.fromisocalendar(year, w, 1)
            month = (monday + dt.timedelta(days=3)).month
            seen[month] += 1
            out[(month, seen[month])] = (w, monday)
        _ISO_CACHE[year] = out
    return _ISO_CACHE[year]


def parse_crosstab(rows):
    """→ (kind, periods, [(labels, values)])"""
    first = next((i for i, v in enumerate(rows[0]) if YEAR_RE.match(v)), None)
    if first is None:
        raise ValueError('연도 헤더 행이 없음')
    for hr, r in enumerate(rows[:6]):
        v = r[first] if len(r) > first else ''
        if DATE_RE.match(v) or WEEK_RE.match(v):
            break
    else:
        raise ValueError('날짜/주차 헤더 행이 없음')
    kind = 'daily' if DATE_RE.match(rows[hr][first]) else 'weekly'

    periods = []
    for c in range(first, len(rows[hr])):
        y = rows[0][c] if c < len(rows[0]) else ''
        v = rows[hr][c]
        m = (DATE_RE if kind == 'daily' else WEEK_RE).match(v)
        if not YEAR_RE.match(y) or not m:
            periods.append(None)
        elif kind == 'daily':
            periods.append(dt.date(int(y), int(m[1]), int(m[2])).isoformat())
        else:
            periods.append((int(y), int(m[1]), int(m[2])))

    if kind == 'weekly' and hr + 1 < len(rows):
        # 진행 중인 주('일마감') — 태블로가 같은 주차 라벨의 지난 연도 열도 그 요일까지만 잘라 내보낸다
        #   (예: 월요일에 내보내면 지난 연도의 같은 주차 열도 월요일 하루치만 담긴다)
        # → 가장 최근 연도의 진행 중 주만 쓰고, 같은 주차 라벨의 다른 연도 열은 버린다(이전 내보내기 · 백업의 완결 값 유지)
        flags = rows[hr + 1]
        part = [periods[c - first] for c in range(first, len(rows[hr]))
                if c < len(flags) and flags[c] == '일마감' and periods[c - first]]
        if part:
            latest, labels = max(p[0] for p in part), {(p[1], p[2]) for p in part}
            periods = [None if p and (p[1], p[2]) in labels and p[0] != latest else p for p in periods]

    records = []
    for r in rows[hr + 1:]:
        cells = r[first:]
        if not any(any(ch.isdigit() for ch in c) for c in cells):
            continue  # 요일 · 주마감 헤더 행
        records.append((r[:first], [to_num(c) for c in cells]))
    return kind, periods, records


MALL_EXISTING_SHARE = 0.7  # 기존회원 거래액 비중이 이보다 크면 MALL 전체 파일 (MALL 은 대부분 기존회원, 첫구매는 훨씬 낮음)


def _overall_is_first_purchase(records):
    """전체관점 파일 내용으로 첫구매/MALL 판별 → (첫구매로 보이는지, 기존회원 거래액 비중)"""
    def total(group):
        return sum(v for labels, vals in records
                   if labels[:3] == ['일평균거래액', group, '*TOTAL'] for v in vals if v)
    tot = total('*TOTAL')
    share = total('3_기존') / tot if tot else None
    return share is not None and share <= MALL_EXISTING_SHARE, share


def put(series, kind, key, period, value):
    if period is None:
        return
    ser = series[kind][key]
    if value is not None or period not in ser:
        ser[period] = value


def infer_md_dates(labels, weekdays=None, today=None):
    """연도 없는 'M/D' 열 → ISO 날짜. 마지막 열을 오늘 이전의 가장 가까운 날짜로 보고 뒤에서부터 연도를 붙인다.
    요일 행이 있으면 대조해 어긋나는 날짜가 있을 때 경고한다."""
    today = today or dt.date.today()
    out, nxt = [None] * len(labels), today + dt.timedelta(days=1)
    for i in range(len(labels) - 1, -1, -1):
        m = DATE_RE.match(labels[i])
        if not m:
            continue
        mo, da = int(m[1]), int(m[2])
        for y in (nxt.year, nxt.year - 1, nxt.year - 2):
            try:
                c = dt.date(y, mo, da)
            except ValueError:  # 윤년이 아닌 해의 2/29
                continue
            if c < nxt:
                break
        out[i] = c.isoformat()
        nxt = c
    if weekdays:
        bad = [d for d, w in zip(out, weekdays)
               if d and w.strip('()') and WEEKDAY_KO[dt.date.fromisoformat(d).weekday()] != w.strip('()')]
        if bad:
            print(f'  [경고] 요일이 맞지 않는 날짜 {len(bad)}개 (예: {bad[:3]}) — 연도 추정을 확인하세요')
    return out


def _load_newmember(rows, series):
    head = rows[0]
    first = next(i for i, h in enumerate(head) if re.match(r'\d{4}\.\s*\d', h))
    col_met, col_age, col_ch = head.index('지표'), head.index('연령대'), head.index('채널')
    dates = []
    for h in head[first:]:
        m = re.match(r'(\d{4})\.\s*(\d{1,2})\.\s*(\d{1,2})', h)
        dates.append(dt.date(int(m[1]), int(m[2]), int(m[3])).isoformat() if m else None)
    n = 0
    for r in rows[1:]:
        base = NEWMEMBER_METRICS.get(r[col_met]) if len(r) > first else None
        if not base or r[col_age] != 'Total':
            continue
        ch = '*TOTAL' if r[col_ch] == 'Total' else r[col_ch]
        for p, v in zip(dates, r[first:]):
            put(series, 'daily', f'{base}|{ch}', p, to_num(v))
        n += 1
    return f'daily 신규회원 당년신규 첫구매 {n}행'


def _load_member_status(rows, series):
    wd = rows[1][1:] if len(rows) > 1 and any(c.startswith('(') for c in rows[1][1:4]) else None
    dates = infer_md_dates(rows[0][1:], wd)
    data = {MEMBER_STATUS[r[0]]: [to_num(v) for v in r[1:]] for r in rows[1:] if r and r[0] in MEMBER_STATUS}
    # 아직 갱신되지 않은 날은 태블로가 0으로 내보낸다(예: 2026-09-14) → 세 값이 모두 0 · 빈 칸인 날은 건너뛴다
    empty = {j for j in range(len(dates)) if all(not (vals[j] if j < len(vals) else None) for vals in data.values())}
    for key, vals in data.items():
        for j, (p, v) in enumerate(zip(dates, vals)):
            if j not in empty:
                put(series, 'daily', f'mem|{key}', p, v)
    note = f' · 값이 모두 0인 날 제외 {[dates[j] for j in sorted(empty)]}' if empty else ''
    return f'daily 회원현황 {len(data)}행 ({dates[0]} ~ {dates[-1]}){note}'


def _load_push(rows, series):
    hr = next(i for i, r in enumerate(rows[:5]) if any(DATE_RE.match(c) for c in r))
    first = next(i for i, c in enumerate(rows[hr]) if DATE_RE.match(c))
    dates = infer_md_dates(rows[hr][first:])
    data = {}
    for r in rows[hr + 1:]:
        seg = PUSH_SEG.get(r[0]) if r else None
        item = PUSH_ITEM.get(r[1]) if len(r) > 1 else None
        if seg and item:
            vals = [to_num(v) for v in r[first:]]
            data[(seg, item)] = vals + [None] * (len(dates) - len(vals))
    blank = [None] * len(dates)
    bad_days = []
    for j, day in enumerate(dates):
        if day is None:
            continue
        stock, chg, add, out = (data.get(('tot', x), blank)[j] for x in ('stock', 'chg', 'add', 'out'))
        if day.endswith('-01-01'):
            # 연초에 전년 '신규' 가 '기존' 으로 재분류되며 추가·이탈에 수만 명이 섞인다 → 추가·이탈만 비운다(재고·증감 유지)
            for (seg, item), vals in data.items():
                if item in ('add', 'out'):
                    vals[j] = None
        elif stock == 0:
            # 원천 오류: 재고가 0으로 떨어지고 전원이 이탈로 잡힘(2025-04-18) → 그날 재고·흐름 모두 제외
            bad_days.append(day)
            for vals in data.values():
                vals[j] = None
        elif (chg is not None and stock and chg == stock) or (add == 0 and out == 0):
            # 원천 오류: 증감·추가 칸에 재고 값이 들어가거나(2025-04-19 · 2026-04-24) 흐름이 모두 0(2026-04-25) → 그날 흐름 제외
            bad_days.append(day)
            for (seg, item), vals in data.items():
                if item != 'stock':
                    vals[j] = None
    for (seg, item), vals in data.items():
        for p, v in zip(dates, vals):
            put(series, 'daily', f'push|{seg}_{item}', p, v)
    note = f' · 흐름 원천 오류로 제외한 날 {bad_days}' if bad_days else ''
    return f'daily 앱푸시 수신동의 {len(data)}행 ({dates[0]} ~ {dates[-1]}){note}'


def _load_app(path, series):
    if path.suffix.lower() == '.csv':
        rows = [tuple(r) for r in read_rows(path)]  # 엑셀을 CSV로 저장해 넣은 경우 (DRM 암호화 회피)
    elif is_drm(path):
        rows = _drm_fallback(path)
    else:
        import openpyxl
        wb = openpyxl.load_workbook(path, read_only=True, data_only=True)
        rows = [tuple(r) for r in wb.worksheets[0].iter_rows(values_only=True)]
        wb.close()
        _save_raw(path, rows)
    header = [str(h or '').strip() for h in rows[0]]
    def as_day(v):
        if isinstance(v, dt.datetime):
            return v.date().isoformat()
        if isinstance(v, dt.date):
            return v.isoformat()
        m = re.match(r'(\d{4})\s*[-./]\s*(\d{1,2})\s*[-./]\s*(\d{1,2})', str(v or ''))
        return dt.date(int(m[1]), int(m[2]), int(m[3])).isoformat() if m else None
    # '날짜' 열이 둘(텍스트 '3월 10일' · 실제 날짜) — 실제 날짜로 읽히는 쪽을 쓴다
    date_col = next((i for i, h in enumerate(header)
                     if h == '날짜' and any(len(r) > i and as_day(r[i]) for r in rows[1:6])), None)
    if date_col is None:
        raise ValueError('날짜 열(YYYY-MM-DD)을 찾지 못함')
    cols = {key: header.index(h) for h, key in APP_COLS.items() if h in header}
    seen, dup, conflict = {}, 0, []
    for r in rows[1:]:
        day = as_day(r[date_col]) if len(r) > date_col else None
        if not day:
            continue
        vals = tuple(r[i] for i in cols.values())
        if day in seen:
            dup += 1
            if seen[day] != vals:
                conflict.append(day)
        seen[day] = vals  # 같은 날짜가 또 나오면 뒤 행을 쓴다
    for day, vals in seen.items():
        for key, v in zip(cols, vals):
            put(series, 'daily', f'app|{key}', day,
                float(v) if isinstance(v, (int, float)) else (to_num(str(v)) if v not in (None, '') else None))
    days = sorted(seen)
    note = ''
    if dup:
        note = f' · 중복 날짜 {dup}건 제거' + (f' (값이 다른 날짜 {conflict[:3]} → 뒤 행 사용)' if conflict else '')
    return f'daily 앱 설치 {len(days)}일 ({days[0]} ~ {days[-1]}){note}'


def load_file(path, series):
    name = path.stem
    if PRODUCT_CSV_KEY in name:
        return '상품관점 → 상품 구성 커버리지(products.js)'  # build_products 에서 읽는다
    if PRODUCT_XLSX_KEY in name:
        return '조직 카테고리별 → 상품 구성(products.js)'  # build_products 에서 읽는다
    if NEWMEMBER_KEY in name:
        return _load_newmember(read_rows(path), series)
    if MEMBER_STATUS_KEY in name:
        return _load_member_status(read_rows(path), series)
    if name.upper().startswith(PUSH_KEY):
        return _load_push(read_rows(path), series)
    if APP_KEY in name:
        return _load_app(path, series)
    kind, periods, records = parse_crosstab(read_rows(path))
    if '전체관점' in name:
        looks_fp, share = _overall_is_first_purchase(records)
        named_fp = '첫구매' in name
        if not (named_fp or looks_fp):
            return None  # MALL 전체 거래액 — 현재 대시보드 미사용
        note = ''
        if named_fp != looks_fp:
            basis = '내용 기준' if looks_fp else '파일명 기준'
            note = f'  ⚠ 파일명과 내용 불일치(기존회원 거래액 비중 {share:.0%}) → {basis} 첫구매로 읽음'
        n = 0
        for labels, vals in records:
            met = OVERALL_METRIC.get(labels[0])
            grp = MEMBER.get(labels[1]) if len(labels) > 2 else None
            if not met or grp is None:
                continue
            for p, v in zip(periods, vals):
                put(series, kind, f'{met}|{grp}|{labels[2]}', p, v)
            n += 1
        return f'{kind} 전체관점(첫구매) {n}행{note}'
    for kw, base in SIMPLE_FILES:
        if kw in name:
            break
    else:
        return None
    for labels, vals in records:
        for p, v in zip(periods, vals):
            put(series, kind, f'{base}|{labels[0]}', p, v)
    return f'{kind} {base} {len(records)}행'


def read_product_xlsx(path):
    """조직 카테고리별 첫구매 실적 xlsx → [(YYYY-MM-DD, bpu, 채널, 카테고리, 브랜드, 상품코드, 상품명, 거래액, 주문고객수)]
    xlsx 읽기가 느려(65만 행 ≈ 45초) 파일 크기·수정시각이 같으면 data/_cache 의 결과를 쓴다."""
    st = path.stat()
    sig = (st.st_size, int(st.st_mtime))
    cache = CACHE_DIR / f'{path.stem}.pkl'
    if cache.exists():
        try:
            cached_sig, rows = pickle.loads(cache.read_bytes())
            if cached_sig == sig:
                return rows
        except Exception:
            pass
    import openpyxl
    wb = openpyxl.load_workbook(path, read_only=True, data_only=True)
    rows = []
    for ws in wb.worksheets:
        it = ws.iter_rows(values_only=True)
        rows.extend(_product_rows(next(it, []), it, f'{path.name}/{ws.title}'))
    wb.close()
    CACHE_DIR.mkdir(parents=True, exist_ok=True)
    cache.write_bytes(pickle.dumps((sig, rows)))
    return rows


def _product_rows(header, it, where):
    """상품 표(헤더 + 행) → [(YYYY-MM-DD, bpu, 채널, 카테고리, 브랜드, 상품코드, 상품명, 거래액, 주문고객수)]
    열은 이름 일부로 찾는다(결제_일자(YYYYMMDD) · AF대분류명 · ADMIN브랜드명 …). xlsx 는 숫자, CSV 는 '1,234' 문자열."""
    header = [str(h or '').strip() for h in header]
    try:
        col = {k: next(i for i, h in enumerate(header) if kw in h) for k, kw in PRODUCT_COLS.items()}
    except StopIteration:
        print(f'  [경고] {where}: 필요한 열을 찾지 못해 건너뜀 {header}')
        return []
    need = max(col.values())

    def num(v):
        if isinstance(v, (int, float)):
            return float(v)
        return to_num(str(v or '')) or 0.0
    out = []
    for r in it:
        if len(r) <= need:
            continue
        ds = str(r[col['date']] or '').replace('-', '')[:8]
        if not ds.isdigit():
            continue

        def text(k, default):
            v = r[col[k]]
            return str(v).strip() if v not in (None, '') else default
        out.append((f'{ds[:4]}-{ds[4:6]}-{ds[6:8]}', text('bpu', '(미지정)'), text('ch', '미분류'),
                    text('cat', '(미지정)'), text('brand', '(미지정)'), text('code', '(코드없음)'), text('name', ''),
                    num(r[col['amt']]), num(r[col['cust']])))
    return out


def read_product_table(path):
    """조직 카테고리별 첫구매 실적 — xlsx(이력) 또는 CSV(일자별 추가분)"""
    if path.suffix.lower() == '.csv':
        rows = read_rows(path)
        return _product_rows(rows[0], rows[1:], path.name) if rows else []
    return read_product_xlsx(path)


def read_product_csv(path):
    """상품관점 일자별 실적 CSV → {(날짜, 채널, BPU): [거래액, 고객수]} — 회원구분·상품군은 *TOTAL 행만"""
    kind, periods, records = parse_crosstab(read_rows(path))
    out = {}
    if kind != 'daily':
        return out

    def total(member):  # 회원구분 member · 채널 · BPU · 상품군 모두 *TOTAL 인 거래액 합
        return sum(v for labels, vals in records
                   if labels[:5] == ['일평균거래액', member, '*TOTAL', '*TOTAL', '*TOTAL'] for v in vals if v)
    whole, existing = total('*TOTAL'), total('3_기존')
    if whole and existing / whole > MALL_EXISTING_SHARE:  # 파일명에 _첫구매 가 없어도 내용으로 판별
        print(f'  [건너뜀] {path.name}: 기존회원 거래액 비중 {existing / whole:.0%} → MALL 전체 파일로 보고 커버리지에서 뺌')
        return out
    for labels, vals in records:
        met = {'일평균거래액': 0, '일평균고객수': 1}.get(labels[0])
        if met is None or len(labels) < 5 or labels[1] != '*TOTAL' or labels[4] != '*TOTAL':
            continue
        for p, v in zip(periods, vals):
            if p is not None and v is not None:
                out.setdefault((p, labels[2], labels[3]), [0.0, 0.0])[met] = v
    return out


def build_products(dirs, last_date, verbose=True):
    """상품 구성 드릴다운 데이터 — 일 × 채널 × 브랜드 경로 × 상품(상위 N, 나머지 -1 = 기타 상품)"""
    by_date, cov_daily = {}, {}
    for d in dirs:
        for f in sorted([*d.glob('*.xlsx'), *d.glob('*.csv')]):
            if PRODUCT_XLSX_KEY not in f.stem or f.name.startswith('~$'):
                continue
            try:
                if f.suffix.lower() == '.xlsx' and is_drm(f):  # 암호화됐으면 마지막으로 읽어 둔 캐시(크기·수정시각과 무관)로 대신 읽는다
                    cache = CACHE_DIR / f'{f.stem}.pkl'
                    if not cache.exists():
                        raise ValueError('DRM 암호화 파일(SCDSA) — 해제 후 넣어 주세요')
                    rows = pickle.loads(cache.read_bytes())[1]
                    print(f'  [경고] {dir_label(d)}/{f.name}: DRM 암호화 → 마지막으로 읽어 둔 캐시로 대신 읽음')
                else:
                    rows = read_product_table(f)
            except Exception as e:  # 깨진 xlsx 하나 때문에 전체 빌드가 멈추지 않게
                print(f'  [오류] {dir_label(d)}/{f.name}: {e}')
                continue
            days = sorted({r[0] for r in rows})
            for day in days:  # 같은 날짜는 나중 파일이 덮어쓴다
                by_date[day] = []
            for r in rows:
                by_date[r[0]].append(r)
            if verbose and days:
                print(f'  읽음   {dir_label(d)}/{f.name}  → 상품 {len(rows):,}행 ({days[0]} ~ {days[-1]})')
        for f in sorted(d.glob('*.csv')):
            if PRODUCT_CSV_KEY in f.stem:
                try:
                    cov_daily.update(read_product_csv(f))
                except Exception as e:
                    print(f'  [오류] {dir_label(d)}/{f.name}: {e}')
    if not by_date:
        return None
    rows = [r for day in sorted(by_date) for r in by_date[day]]
    last = dt.date.fromisoformat(last_date or max(by_date))
    start = dt.date.fromisoformat(min(by_date))
    day_of = {}

    def day_idx(day):  # 시작일부터 며칠째 — 대시보드가 이 값으로 월·주·일 기간을 다시 모은다
        if day not in day_of:
            day_of[day] = (dt.date.fromisoformat(day) - start).days
        return day_of[day]

    known = [c for c in CHANNEL_ORDER if c != '*TOTAL']
    chs = known + sorted({r[2] for r in rows} - set(known))
    bpus, cats, brands = (sorted({r[i] for r in rows}) for i in (1, 3, 4))
    ci, bi, ki, ri = ({v: i for i, v in enumerate(x)} for x in (chs, bpus, cats, brands))
    paths, path_idx, names, ytot = [], {}, {}, defaultdict(float)
    for r in rows:
        key = (bi[r[1]], ki[r[3]], ri[r[4]])
        if key not in path_idx:
            path_idx[key] = len(paths) // 3
            paths.extend(key)
        if r[6]:
            names[r[5]] = r[6]  # 날짜순으로 돌므로 마지막 값 = 최신 상품명
        ytot[(r[0][:4], path_idx[key], r[5])] += r[7]
    by_path = defaultdict(list)
    for (y, p, code), a in ytot.items():
        by_path[(y, p)].append((a, code))
    keep, prods, prod_idx = set(), [], {}
    for (y, p), lst in sorted(by_path.items()):
        for a, code in sorted(lst, reverse=True)[:PRODUCT_TOP_N]:
            keep.add((y, p, code))
            if code not in prod_idx:
                prod_idx[code] = len(prods) // 2
                prods.extend([code, names.get(code, '')])

    facts = defaultdict(lambda: [0.0, 0.0])
    for r in rows:
        p = path_idx[(bi[r[1]], ki[r[3]], ri[r[4]])]
        q = prod_idx[r[5]] if (r[0][:4], p, r[5]) in keep else -1
        o = facts[(day_idx(r[0]), ci[r[2]], p, q)]
        o[0] += r[7]
        o[1] += r[8]
    covm = defaultdict(lambda: [0.0, 0.0])
    for (day, ch, bpu), (a, u) in cov_daily.items():
        if ch != '*TOTAL' and ch not in ci:
            ci[ch] = len(chs)
            chs.append(ch)
        if bpu != '*TOTAL' and bpu not in bi:
            bi[bpu] = len(bpus)
            bpus.append(bpu)
        if day < start.isoformat():
            continue
        o = covm[(day_idx(day), -1 if ch == '*TOTAL' else ci[ch], -1 if bpu == '*TOTAL' else bi[bpu])]
        o[0] += a
        o[1] += u

    # 날짜순으로 정렬하고 날짜는 앞 행과의 차이만 저장(대부분 0) → 파일 크기 절약
    flat, prev_d = [], 0
    for (d, c, p, q), (a, u) in sorted(facts.items()):
        if round(a) == 0 and round(u) == 0:
            continue
        flat.extend([d - prev_d, c, p, q, round(a), round(u)])
        prev_d = d
    cov, prev_d = [], 0
    for (d, c, b), (a, u) in sorted(covm.items()):
        cov.extend([d - prev_d, c, b, round(a), round(u)])
        prev_d = d
    days = max([k[0] for k in facts] + [k[0] for k in covm]) + 1
    return {
        'meta': {'built': dt.datetime.now().strftime('%Y-%m-%d %H:%M'), 'topN': PRODUCT_TOP_N, 'start': start.isoformat(),
                 'days': days, 'lastDate': last.isoformat(), 'rows': len(rows), 'products': len(names),
                 'encoding': 'f=[날짜차분, 채널, 경로, 상품(-1=기타), 거래액, 주문고객수] · cov=[날짜차분, 채널(-1=전체), BPU(-1=전체), 거래액, 고객수]'},
        'ch': chs, 'bpu': bpus, 'cat': cats, 'brand': brands, 'paths': paths, 'prods': prods, 'f': flat, 'cov': cov,
    }


def _raw_dirs(base):
    """base 와 그 아래 모든 폴더 중 CSV 가 든 폴더 — 예: 9월2주차 · 2026년/0914 · 2026년/0915 …"""
    out = []
    for d in [base, *sorted(p for p in base.rglob('*') if p.is_dir())]:
        rel = d.relative_to(ROOT).parts if d.is_relative_to(ROOT) else d.parts
        if any(part in SKIP_DIRS or part.startswith('.') for part in rel):
            continue
        if any(d.glob('*.csv')):
            out.append(d)
    return out


def dir_label(d):
    return d.relative_to(ROOT).as_posix() if d.is_relative_to(ROOT) else d.name


def discover(args):
    bases = [Path(a) if Path(a).is_absolute() else ROOT / a for a in args] if args else [ROOT]
    dirs = list(dict.fromkeys(d for b in bases if b.is_dir() for d in _raw_dirs(b)))
    # 같은 날짜 값은 나중 폴더가 덮어쓴다 — 가장 최근에 내보낸 CSV 가 있는 폴더를 나중에
    return sorted(dirs, key=lambda d: max(f.stat().st_mtime for f in d.glob('*.csv')))


def _round(key, v):
    if v is None:
        return None
    base = key.split('|')[0]
    if base in ('amt', 'nya', 'nyna'):
        return round(v)
    if base in ('crn', 'fpn', 'cust'):
        return round(v, 3)
    return round(v, 2)


def _emit(src, keys, days=None, rename=None):
    out = {}
    for key, ser in src.items():
        base = key.split('|')[0]
        if base in ('cr', 'fpr'):
            continue
        vals = [ser.get(k) for k in keys]
        if days and base in ('amt', 'cust'):  # 주별 원천은 일평균 → 주 합계
            vals = [None if v is None else v * days[k] for v, k in zip(vals, keys)]
        out[key] = vals
    for rate, den, name in RATE_TO_COUNT:
        for key, ser in src.items():
            if not key.startswith(rate + '|'):
                continue
            ch = key.split('|', 1)[1]
            dser = src.get(f'{den}|{ch}', {})
            out[f'{name}|{ch}'] = [None if ser.get(k) is None or dser.get(k) is None
                                   else ser[k] * dser[k] for k in keys]
    rename = rename or {}

    def name(k):
        base, rest = k.split('|', 1)
        return f'{rename.get(base, base)}|{rest}'
    return {name(k): [_round(k, v) for v in arr] for k, arr in sorted(out.items())
            if any(v is not None for v in arr)}


def build(dirs, verbose=True):
    series = {'daily': defaultdict(dict), 'weekly': defaultdict(dict)}
    for d in dirs:
        app_xlsx = [x for x in d.glob('*.xlsx') if APP_KEY in x.stem and not x.name.startswith('~$')]
        for f in sorted(list(d.glob('*.csv')) + app_xlsx):
            try:
                res = load_file(f, series)
            except Exception as e:  # 깨진 파일 하나 때문에 전체가 멈추지 않게
                print(f'  [오류] {dir_label(d)}/{f.name}: {e}')
                continue
            if verbose:
                print(f'  {"읽음  " if res else "건너뜀"} {dir_label(d)}/{f.name}' + (f'  → {res}' if res else ''))

    daily, weekly = series['daily'], series['weekly']
    dates = sorted({p for ser in daily.values() for p, v in ser.items() if v is not None})
    last = dt.date.fromisoformat(dates[-1]) if dates else None

    weeks, wkeys = [], []
    for y, m, n in sorted({p for ser in weekly.values() for p, v in ser.items() if v is not None}):
        cal = iso_week_map(y)
        if (m, n) not in cal:
            print(f'  [경고] {y}년 {m}월 {n}주차는 ISO 달력에 없음 — 제외')
            continue
        w, monday = cal[(m, n)]
        d = 7
        if last and monday <= last <= monday + dt.timedelta(days=6):
            d = (last - monday).days + 1  # 진행 중인 주
        weeks.append({'y': y, 'm': m, 'n': n, 'w': w, 'start': monday.isoformat(), 'd': d})
        wkeys.append((y, m, n))
    days_of = {k: p['d'] for k, p in zip(wkeys, weeks)}

    chans = {k.split('|')[-1] for src in (daily, weekly) for k in src if k.split('|')[0] in CHANNEL_BASES}
    channels = [c for c in CHANNEL_ORDER if c in chans] + sorted(chans - set(CHANNEL_ORDER))

    return {
        'meta': {
            'built': dt.datetime.now().strftime('%Y-%m-%d %H:%M'),
            'sources': [dir_label(d) for d in dirs],
            'lastDate': dates[-1] if dates else None,
            'channels': channels,
        },
        'daily': {'p': dates, 's': _emit(daily, dates)},
        'weekly': {'p': weeks, 's': _emit(weekly, wkeys, days_of, rename={'tr': 'tru', 'fpn': 'fpnu'})},
    }


def read_kpi():
    """kpi.js(자바스크립트 객체) → dict — 백업 파일에 KPI 목표도 담으려고 읽는다. 못 읽으면 None"""
    try:
        text = KPI_JS.read_text(encoding='utf-8')
        body = text[text.index('{'):text.rindex('}') + 1]
    except (OSError, ValueError):
        return None
    body = re.sub(r'//[^\n]*', '', body)                                                      # 주석
    body = re.sub(r"'([^'\\]*)'", lambda m: json.dumps(m[1], ensure_ascii=False), body)       # '문자열' → "문자열"
    body = re.sub(r'([{,]\s*)([A-Za-z_]\w*)\s*:', r'\1"\2":', body)                           # 키에 따옴표
    body = re.sub(r',(\s*[}\]])', r'\1', body)                                                # 끝 쉼표
    try:
        return json.loads(body)
    except ValueError as e:
        print(f'  [경고] kpi.js 를 읽지 못해 백업에 KPI를 넣지 않음: {e}')
        return None


def write_backup(data, prod, path=None):
    """대시보드 '백업 다운로드'와 같은 형식(.json.gz) — '백업 불러오기'로 바로 열 수 있다"""
    import gzip
    payload = {'format': BACKUP_FORMAT, 'version': 1, 'created': dt.datetime.now().astimezone().isoformat(timespec='seconds'),
               'source': 'build_data.py', 'lastDate': data['meta']['lastDate'], 'built': data['meta']['built'],
               'data': data, 'kpi': read_kpi(), 'prod': prod}
    path = path or BACKUP_DIR / f"첫구매대시보드_백업_데이터{(data['meta']['lastDate'] or '').replace('-', '')}.json.gz"
    path.parent.mkdir(parents=True, exist_ok=True)
    path.write_bytes(gzip.compress(json.dumps(payload, ensure_ascii=False, separators=(',', ':')).encode('utf-8'), mtime=0))
    return path


def main(argv):
    dirs = discover(argv)
    if not dirs:
        sys.exit('raw 폴더를 찾지 못했습니다 (CSV가 든 하위 폴더가 필요)')
    print('raw 폴더:', ', '.join(dir_label(d) for d in dirs))
    data = build(dirs)
    OUT.parent.mkdir(exist_ok=True)
    OUT.write_text('window.DASH_DATA = ' + json.dumps(data, ensure_ascii=False, separators=(',', ':')) + ';\n',
                   encoding='utf-8')
    d, w = data['daily'], data['weekly']
    print(f'\n일별 {len(d["p"])}일 ({d["p"][0] if d["p"] else "-"} ~ {data["meta"]["lastDate"]}) · 시리즈 {len(d["s"])}')
    if w['p']:
        print(f'주별 {len(w["p"])}주 ({w["p"][0]["y"]}-{w["p"][0]["m"]}월{w["p"][0]["n"]}주 ~ '
              f'{w["p"][-1]["y"]}-{w["p"][-1]["m"]}월{w["p"][-1]["n"]}주) · 시리즈 {len(w["s"])}')
    print(f'→ {OUT} ({OUT.stat().st_size / 1024:.0f} KB)')

    print('\n상품 구성:')
    prod = build_products(dirs, data['meta']['lastDate'])
    if prod:
        OUT_PROD.write_text('window.DASH_PROD = ' + json.dumps(prod, ensure_ascii=False, separators=(',', ':')) + ';\n',
                            encoding='utf-8')
        print(f'상품 원천 {prod["meta"]["rows"]:,}행 → 집계 {len(prod["f"]) // 6:,}건 · 개별 상품 {len(prod["prods"]) // 2:,}개'
              f' (브랜드별 연간 상위 {PRODUCT_TOP_N}) → {OUT_PROD} ({OUT_PROD.stat().st_size / 1024 / 1024:.1f} MB)')
    else:
        print('  상품 원천 없음 — 상품 구성 섹션은 숨겨진다')
    bk = write_backup(data, prod)
    print(f'\n백업 → {bk} ({bk.stat().st_size / 1024 / 1024:.1f} MB) · 대시보드 \'백업 불러오기\'로 열 수 있음')


if __name__ == '__main__':
    main(sys.argv[1:])
